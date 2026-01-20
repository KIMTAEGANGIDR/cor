#!/usr/bin/env python3
"""Run OCR-vs-extracted-text comparison on OCR candidate PDFs.

This focuses on higher-risk documents by sampling from an OCR candidate list,
and compares image OCR text against embedded PDF text on 3 pages per document.

Usage:
  .venv/bin/python scripts/multimodal_candidate_report.py \
    --candidates docs/reports/ocr_candidates_full.json \
    --count 100 --seed 42 --output-dir docs/reports
"""

from __future__ import annotations

import argparse
import json
import random
import re
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path
from statistics import mean
from typing import Iterable

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def load_candidates(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("candidates", [])


def sample_candidates(items: list[dict], count: int, seed: int) -> list[dict]:
    if count <= 0 or count >= len(items):
        return items
    random.seed(seed)
    return random.sample(items, count)


def pick_pages(page_count: int) -> list[int]:
    if page_count <= 0:
        return []
    mid = page_count // 2
    idxs = {0, mid, page_count - 1}
    return sorted(i for i in idxs if 0 <= i < page_count)


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def tokenize(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def jaccard(a: Iterable[str], b: Iterable[str]) -> float:
    set_a = set(a)
    set_b = set(b)
    if not set_a and not set_b:
        return 1.0
    if not set_a or not set_b:
        return 0.0
    return len(set_a & set_b) / len(set_a | set_b)


def alpha_ratio(text: str) -> float:
    raw = re.sub(r"\s+", "", text)
    if not raw:
        return 0.0
    good = sum(1 for ch in raw if ch.isalnum())
    return good / len(raw)


def sample_text(text: str, limit: int = 180) -> str:
    t = normalize_text(text)
    return t[:limit]


def run_tesseract(img_path: Path, lang: str, psm: int) -> str:
    cmd = [
        "tesseract",
        str(img_path),
        "stdout",
        "-l",
        lang,
        "--psm",
        str(psm),
    ]
    out = subprocess.check_output(cmd, stderr=subprocess.DEVNULL)
    return out.decode("utf-8", errors="ignore")


def ocr_page(page: fitz.Page, dpi: int, lang: str, psm: int, temp_dir: Path) -> str:
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csGRAY)
    img_path = temp_dir / "page.png"
    pix.save(img_path)
    return run_tesseract(img_path, lang, psm)


def analyze_pdf(
    pdf_path: Path,
    dpi: int,
    lang: str,
    psm: int,
    thresholds: dict,
) -> dict:
    result = {
        "path": str(pdf_path),
        "type": pdf_path.parent.name,
        "pages": 0,
        "page_metrics": [],
        "flags": [],
        "avg_similarity": 0.0,
        "worst_page": None,
        "error": None,
    }

    try:
        with fitz.open(pdf_path) as doc:
            page_count = len(doc)
            result["pages"] = page_count
            page_indexes = pick_pages(page_count)
            if not page_indexes:
                result["error"] = "empty_pdf"
                return result

            similarities = []
            low_similarity_pages = 0
            low_text_pages = 0
            garbage_pages = 0
            ocr_rich_pages = 0

            with tempfile.TemporaryDirectory() as tmp_dir:
                temp_dir = Path(tmp_dir)
                for page_idx in page_indexes:
                    page = doc.load_page(page_idx)
                    extracted = page.get_text() or ""
                    ocr_text = ""
                    ocr_error = None
                    try:
                        ocr_text = ocr_page(page, dpi, lang, psm, temp_dir)
                    except Exception as exc:
                        ocr_error = str(exc)

                    extracted_tokens = tokenize(extracted)
                    ocr_tokens = tokenize(ocr_text)
                    sim = jaccard(extracted_tokens, ocr_tokens)
                    similarities.append(sim)

                    extracted_len = len(extracted)
                    ocr_len = len(ocr_text)
                    alpha = alpha_ratio(extracted)

                    flags = []
                    if ocr_error:
                        flags.append("ocr_error")
                    if extracted_len < thresholds["min_extracted_len"] and ocr_len >= thresholds["min_ocr_len"]:
                        flags.append("low_extracted_text")
                        low_text_pages += 1
                    if sim < thresholds["min_similarity"] and len(ocr_tokens) >= thresholds["min_ocr_tokens"]:
                        flags.append("low_similarity")
                        low_similarity_pages += 1
                    if alpha < thresholds["min_alpha_ratio"] and extracted_len >= thresholds["min_extracted_len"]:
                        flags.append("garbage_text")
                        garbage_pages += 1
                    if len(ocr_tokens) >= thresholds["min_ocr_tokens"] and len(extracted_tokens) <= thresholds["max_extracted_tokens"]:
                        flags.append("ocr_rich_text")
                        ocr_rich_pages += 1

                    metric = {
                        "page": page_idx + 1,
                        "extracted_len": extracted_len,
                        "ocr_len": ocr_len,
                        "extracted_tokens": len(extracted_tokens),
                        "ocr_tokens": len(ocr_tokens),
                        "similarity": round(sim, 3),
                        "alpha_ratio": round(alpha, 3),
                        "flags": flags,
                        "extracted_sample": sample_text(extracted),
                        "ocr_sample": sample_text(ocr_text),
                    }
                    result["page_metrics"].append(metric)

            result["avg_similarity"] = round(mean(similarities), 3) if similarities else 0.0
            worst = min(result["page_metrics"], key=lambda m: m["similarity"]) if result["page_metrics"] else None
            result["worst_page"] = worst

            if low_similarity_pages > 0:
                result["flags"].append("low_similarity_doc")
            if low_text_pages > 0:
                result["flags"].append("low_extracted_text_doc")
            if garbage_pages > 0:
                result["flags"].append("garbage_text_doc")
            if ocr_rich_pages > 0:
                result["flags"].append("ocr_rich_text_doc")

    except Exception as exc:
        result["error"] = str(exc)

    return result


def build_docx_report(summary: dict, rows: list[dict], output_path: Path) -> None:
    doc = Document()
    title = doc.add_heading("OCR 후보군 100건 이미지-텍스트 비교 보고서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("설명: OCR 후보군에서 100건을 무작위로 추출했습니다.")
    doc.add_paragraph("방법: 각 문서의 1페이지/중간/마지막 페이지를 이미지 OCR로 읽고,")
    doc.add_paragraph("      같은 페이지의 PDF 내장 텍스트와 비교했습니다.")

    doc.add_heading("1. 핵심 요약", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "값"
    summary_rows = [
        ("총 샘플 수", summary.get("files_total")),
        ("분석 성공", summary.get("files_ok")),
        ("분석 오류", summary.get("files_error")),
        ("평균 유사도(3페이지 평균)", summary.get("avg_similarity")),
        ("유사도 낮은 문서", summary.get("low_similarity_docs")),
        ("OCR 텍스트 풍부/추출 텍스트 빈약", summary.get("ocr_rich_text_docs")),
        ("추출 텍스트 깨짐 의심", summary.get("garbage_text_docs")),
    ]
    for label, value in summary_rows:
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(value)

    doc.add_heading("2. 해석", level=1)
    for line in summary.get("interpretation", []):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("3. 유형별 분포", level=1)
    type_breakdown = summary.get("type_breakdown", {})
    if type_breakdown:
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "유형"
        table.rows[0].cells[1].text = "건수"
        table.rows[0].cells[2].text = "유사도 낮음"
        table.rows[0].cells[3].text = "OCR 풍부/추출 빈약"
        for key, stats in type_breakdown.items():
            row = table.add_row().cells
            row[0].text = key
            row[1].text = str(stats.get("count", 0))
            row[2].text = str(stats.get("low_similarity", 0))
            row[3].text = str(stats.get("ocr_rich", 0))
    else:
        doc.add_paragraph("유형별 통계가 없습니다.")

    doc.add_heading("4. 의심 사례 (유사도 최저 10건)", level=1)
    for row in summary.get("worst_cases", []):
        doc.add_paragraph(row["path"], style="List Bullet")
        doc.add_paragraph(
            f"  avg_similarity={row['avg_similarity']} | flags={', '.join(row['flags'])}"
        )
        doc.add_paragraph(
            f"  worst_page={row['page']} | sim={row['similarity']} | "
            f"extracted_len={row['extracted_len']} | ocr_len={row['ocr_len']}"
        )
        doc.add_paragraph(f"  extracted_sample: {row['extracted_sample']}")
        doc.add_paragraph(f"  ocr_sample: {row['ocr_sample']}")

    doc.add_heading("5. 기준값(임계치)", level=1)
    thresholds = summary.get("thresholds", {})
    for k, v in thresholds.items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="OCR candidate multimodal report")
    parser.add_argument("--candidates", required=True, help="OCR candidates JSON")
    parser.add_argument("--count", type=int, default=100, help="Sample size")
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument("--dpi", type=int, default=200, help="Render DPI for OCR")
    parser.add_argument("--lang", default="ind+eng", help="Tesseract languages")
    parser.add_argument("--psm", type=int, default=6, help="Tesseract PSM")
    parser.add_argument("--output-dir", default="docs/reports", help="Output directory")
    args = parser.parse_args()

    candidates_path = Path(args.candidates)
    items = load_candidates(candidates_path)
    if not items:
        print("No candidates found.")
        return 1

    sample = sample_candidates(items, args.count, args.seed)

    thresholds = {
        "min_similarity": 0.25,
        "min_ocr_tokens": 40,
        "max_extracted_tokens": 20,
        "min_extracted_len": 200,
        "min_ocr_len": 200,
        "min_alpha_ratio": 0.6,
    }

    rows = []
    for item in sample:
        pdf_path = Path(item["path"])
        rows.append(analyze_pdf(pdf_path, args.dpi, args.lang, args.psm, thresholds))

    ok_rows = [r for r in rows if not r.get("error")]
    error_rows = [r for r in rows if r.get("error")]

    summary = {
        "files_total": len(rows),
        "files_ok": len(ok_rows),
        "files_error": len(error_rows),
        "avg_similarity": round(mean(r["avg_similarity"] for r in ok_rows), 3) if ok_rows else 0.0,
        "low_similarity_docs": sum(1 for r in ok_rows if "low_similarity_doc" in r.get("flags", [])),
        "ocr_rich_text_docs": sum(1 for r in ok_rows if "ocr_rich_text_doc" in r.get("flags", [])),
        "garbage_text_docs": sum(1 for r in ok_rows if "garbage_text_doc" in r.get("flags", [])),
        "thresholds": thresholds,
    }

    interpretation = []
    if summary["low_similarity_docs"]:
        interpretation.append(
            "OCR 텍스트와 PDF 내장 텍스트가 크게 다른 문서가 다수 확인되었습니다."
        )
    if summary["ocr_rich_text_docs"]:
        interpretation.append(
            "이미지에는 글자가 충분히 보이지만, PDF 내장 텍스트가 거의 없는 경우가 있습니다."
        )
    if summary["garbage_text_docs"]:
        interpretation.append(
            "PDF 내장 텍스트가 깨져 있어 구조화 전에 정제/재추출이 필요합니다."
        )
    if not interpretation:
        interpretation.append(
            "OCR 후보군에서도 큰 불일치는 많지 않았습니다. 표본 확대가 필요합니다."
        )
    summary["interpretation"] = interpretation

    by_type = defaultdict(list)
    for row in ok_rows:
        by_type[row["type"]].append(row)

    summary["type_breakdown"] = {
        k: {
            "count": len(v),
            "low_similarity": sum(1 for r in v if "low_similarity_doc" in r.get("flags", [])),
            "ocr_rich": sum(1 for r in v if "ocr_rich_text_doc" in r.get("flags", [])),
        }
        for k, v in sorted(by_type.items(), key=lambda x: len(x[1]), reverse=True)
    }

    worst_cases = []
    for row in sorted(ok_rows, key=lambda r: r.get("avg_similarity", 1.0))[:10]:
        wp = row.get("worst_page") or {}
        worst_cases.append({
            "path": row["path"],
            "avg_similarity": row.get("avg_similarity"),
            "flags": row.get("flags", []),
            "page": wp.get("page"),
            "similarity": wp.get("similarity"),
            "extracted_len": wp.get("extracted_len"),
            "ocr_len": wp.get("ocr_len"),
            "extracted_sample": wp.get("extracted_sample"),
            "ocr_sample": wp.get("ocr_sample"),
        })
    summary["worst_cases"] = worst_cases

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / f"multimodal_candidate_{args.count}.json"
    docx_path = output_dir / f"multimodal_candidate_{args.count}.docx"

    with json_path.open("w", encoding="utf-8") as f:
        json.dump({"summary": summary, "rows": rows}, f, ensure_ascii=False, indent=2)

    build_docx_report(summary, rows, docx_path)

    print(f"Wrote: {json_path}")
    print(f"Wrote: {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
