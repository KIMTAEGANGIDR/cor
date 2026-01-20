#!/usr/bin/env python3
"""Compare OCR noise removal before/after on a random sample.

Usage:
  .venv/bin/python scripts/ocr_noise_compare.py --count 300 --output-dir docs/reports
"""

from __future__ import annotations

import argparse
import json
import random
import re
import subprocess
import tempfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path
from statistics import mean
from typing import Iterable

import fitz  # PyMuPDF

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except Exception:
    DOCX_AVAILABLE = False

from peraturan.src.services.ocr_cleaner import clean_ocr_pages, extract_page_lines, find_repeated_lines


def pick_pages(page_count: int, mode: str) -> list[int]:
    if page_count <= 0:
        return []
    if mode == "first":
        return [0]
    mid = page_count // 2
    idxs = {0, mid, page_count - 1}
    return sorted(i for i in idxs if 0 <= i < page_count)


def tokenize(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def jaccard(a: Iterable[str], b: Iterable[str]) -> float:
    set_a = set(a)
    set_b = set(b)
    if not set_a and not set_b:
        return 1.0
    if not set_a or not set_b:
        return 0.0
    return len(set_a & set_b) / len(set_a | set_b)


def detect_markers(text: str) -> dict:
    t = text.lower()
    return {
        "memutuskan": "memutuskan" in t,
        "menetapkan": "menetapkan" in t,
        "pasal": "pasal" in t,
        "bab": bool(re.search(r"\bbab\s+[ivxlcdm]+", t)),
    }


def run_tesseract(img_path: Path, lang: str, psm: int) -> str:
    cmd = [
        "tesseract",
        str(img_path),
        "stdout",
        "-l",
        lang,
        "--psm",
        str(psm),
    ]
    out = subprocess.check_output(cmd, stderr=subprocess.DEVNULL)
    return out.decode("utf-8", errors="ignore")


def ocr_page(page: fitz.Page, dpi: int, lang: str, psm: int, temp_dir: Path) -> str:
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csGRAY)
    img_path = temp_dir / "page.png"
    pix.save(img_path)
    return run_tesseract(img_path, lang, psm)


def analyze_pdf(
    pdf_path: Path,
    dpi: int,
    lang: str,
    psm: int,
    page_mode: str,
    min_embed_tokens: int,
) -> dict:
    result = {
        "path": str(pdf_path),
        "type": pdf_path.parent.name,
        "pages": 0,
        "baseline": {},
        "cleaned": {},
        "delta": {},
        "error": None,
    }

    try:
        with fitz.open(pdf_path) as doc:
            page_count = len(doc)
            result["pages"] = page_count
            page_indexes = pick_pages(page_count, page_mode)
            if not page_indexes:
                result["error"] = "empty_pdf"
                return result

            page_texts = []
            embed_texts = []

            with tempfile.TemporaryDirectory() as tmp_dir:
                temp_dir = Path(tmp_dir)
                for page_idx in page_indexes:
                    page = doc.load_page(page_idx)
                    embed_texts.append(page.get_text() or "")
                    ocr_text = ocr_page(page, dpi, lang, psm, temp_dir)
                    page_texts.append(ocr_text)

            raw_text = "\n".join(page_texts)
            embed_text = "\n".join(embed_texts)
            embed_tokens = tokenize(embed_text)

            page_lines = [extract_page_lines(text) for text in page_texts]
            repeated_lines, repeated_ratio = find_repeated_lines(page_lines, len(page_texts))

            baseline_markers = detect_markers(raw_text)
            baseline_similarity = None
            if len(embed_tokens) >= min_embed_tokens:
                baseline_similarity = round(jaccard(tokenize(raw_text), embed_tokens), 3)

            baseline = {
                "chars": len(raw_text),
                "lines": sum(len(t.splitlines()) for t in page_texts),
                "repeated_line_ratio": round(repeated_ratio, 3),
                "markers": baseline_markers,
                "similarity": baseline_similarity,
            }

            cleaned = clean_ocr_pages(page_texts)
            cleaned_text = "\n".join(cleaned.cleaned_pages)
            cleaned_lines = [extract_page_lines(text) for text in cleaned.cleaned_pages]
            _, cleaned_ratio = find_repeated_lines(cleaned_lines, len(cleaned.cleaned_pages))

            cleaned_markers = detect_markers(cleaned_text)
            cleaned_similarity = None
            if len(embed_tokens) >= min_embed_tokens:
                cleaned_similarity = round(jaccard(tokenize(cleaned_text), embed_tokens), 3)

            cleaned_metrics = {
                "chars": len(cleaned_text),
                "lines": sum(len(t.splitlines()) for t in cleaned.cleaned_pages),
                "repeated_line_ratio": round(cleaned_ratio, 3),
                "markers": cleaned_markers,
                "similarity": cleaned_similarity,
                "removed_lines": cleaned.removed_lines,
                "total_lines": cleaned.total_lines,
            }

            delta_similarity = None
            if baseline_similarity is not None and cleaned_similarity is not None:
                delta_similarity = round(cleaned_similarity - baseline_similarity, 3)

            result["baseline"] = baseline
            result["cleaned"] = cleaned_metrics
            result["delta"] = {
                "similarity": delta_similarity,
                "repeated_line_ratio": round(baseline["repeated_line_ratio"] - cleaned_metrics["repeated_line_ratio"], 3),
            }

    except Exception as exc:
        result["error"] = str(exc)

    return result


def analyze_pdf_worker(args: tuple[Path, int, str, int, str, int]) -> dict:
    pdf_path, dpi, lang, psm, page_mode, min_embed_tokens = args
    return analyze_pdf(pdf_path, dpi, lang, psm, page_mode, min_embed_tokens)


def summarize(rows: list[dict]) -> dict:
    ok_rows = [r for r in rows if not r.get("error")]
    before_ratios = [r["baseline"]["repeated_line_ratio"] for r in ok_rows]
    after_ratios = [r["cleaned"]["repeated_line_ratio"] for r in ok_rows]
    removed_ratios = [
        (r["cleaned"]["removed_lines"] / r["cleaned"]["total_lines"])
        if r["cleaned"]["total_lines"] else 0.0
        for r in ok_rows
    ]

    sim_before = [r["baseline"]["similarity"] for r in ok_rows if r["baseline"]["similarity"] is not None]
    sim_after = [r["cleaned"]["similarity"] for r in ok_rows if r["cleaned"]["similarity"] is not None]
    sim_delta = [
        r["delta"]["similarity"] for r in ok_rows if r["delta"]["similarity"] is not None
    ]

    def marker_rate(key: str, phase: str) -> float:
        vals = [r[phase]["markers"][key] for r in ok_rows]
        return round(sum(1 for v in vals if v) / len(vals), 3) if vals else 0.0

    summary = {
        "files_total": len(rows),
        "files_ok": len(ok_rows),
        "files_error": len(rows) - len(ok_rows),
        "avg_repeated_ratio_before": round(mean(before_ratios), 3) if before_ratios else 0.0,
        "avg_repeated_ratio_after": round(mean(after_ratios), 3) if after_ratios else 0.0,
        "avg_removed_line_ratio": round(mean(removed_ratios), 3) if removed_ratios else 0.0,
        "avg_similarity_before": round(mean(sim_before), 3) if sim_before else None,
        "avg_similarity_after": round(mean(sim_after), 3) if sim_after else None,
        "avg_similarity_delta": round(mean(sim_delta), 3) if sim_delta else None,
        "improved_similarity_docs": sum(1 for v in sim_delta if v and v > 0),
        "marker_rate_before": {
            "pasal": marker_rate("pasal", "baseline"),
            "bab": marker_rate("bab", "baseline"),
            "memutuskan": marker_rate("memutuskan", "baseline"),
            "menetapkan": marker_rate("menetapkan", "baseline"),
        },
        "marker_rate_after": {
            "pasal": marker_rate("pasal", "cleaned"),
            "bab": marker_rate("bab", "cleaned"),
            "memutuskan": marker_rate("memutuskan", "cleaned"),
            "menetapkan": marker_rate("menetapkan", "cleaned"),
        },
    }

    return summary


def build_docx_report(summary: dict, rows: list[dict], output_path: Path) -> None:
    if not DOCX_AVAILABLE:
        return

    doc = Document()
    title = doc.add_heading("OCR 노이즈 제거 전/후 비교 보고서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("목적: OCR 결과에서 헤더/푸터 등 반복 노이즈를 제거했을 때 품질이 개선되는지 확인합니다.")
    doc.add_paragraph("방법: 각 문서의 1/중간/마지막 페이지를 OCR 후, 반복 라인 및 고정 노이즈 패턴을 제거합니다.")

    doc.add_heading("1. 핵심 요약", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "값"
    for key, value in [
        ("총 샘플 수", summary.get("files_total")),
        ("분석 성공", summary.get("files_ok")),
        ("분석 오류", summary.get("files_error")),
        ("반복 라인 비율(전)", summary.get("avg_repeated_ratio_before")),
        ("반복 라인 비율(후)", summary.get("avg_repeated_ratio_after")),
        ("평균 삭제 라인 비율", summary.get("avg_removed_line_ratio")),
        ("유사도(전)", summary.get("avg_similarity_before")),
        ("유사도(후)", summary.get("avg_similarity_after")),
        ("유사도 개선 문서 수", summary.get("improved_similarity_docs")),
    ]:
        row = table.add_row().cells
        row[0].text = str(key)
        row[1].text = str(value)

    doc.add_heading("2. 구조 마커 검출률(전/후)", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "마커"
    table.rows[0].cells[1].text = "전"
    table.rows[0].cells[2].text = "후"
    for key in ("pasal", "bab", "memutuskan", "menetapkan"):
        row = table.add_row().cells
        row[0].text = key
        row[1].text = str(summary.get("marker_rate_before", {}).get(key))
        row[2].text = str(summary.get("marker_rate_after", {}).get(key))

    doc.add_heading("3. 해석", level=1)
    interpretations = [
        "반복 라인 비율이 감소하면 헤더/푸터 노이즈가 제거되고 있다는 의미입니다.",
        "유사도가 개선된 문서가 많을수록 정제 효과가 크다고 볼 수 있습니다.",
        "구조 마커(Pasal, BAB 등)가 더 잘 검출되면 본문 구조화 성공률이 높아집니다.",
    ]
    for line in interpretations:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("4. 다음 조치", level=1)
    actions = [
        "OCR 파이프라인에 반복 라인 제거를 기본 단계로 포함",
        "정제 전/후 유사도 차이를 기준으로 자동 재처리 규칙 설정",
        "노이즈 패턴 사전을 유형별로 확장",
    ]
    for line in actions:
        doc.add_paragraph(line, style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="OCR noise compare")
    parser.add_argument("--root", default="peraturan/data/pdfs", help="PDF root directory")
    parser.add_argument("--count", type=int, default=300, help="Sample size")
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument("--dpi", type=int, default=150, help="Render DPI for OCR")
    parser.add_argument("--lang", default="ind+eng", help="Tesseract languages")
    parser.add_argument("--psm", type=int, default=6, help="Tesseract PSM")
    parser.add_argument("--page-mode", choices=["first", "three"], default="three", help="Pages to OCR")
    parser.add_argument("--min-embed-tokens", type=int, default=40, help="Min embedded tokens for similarity")
    parser.add_argument("--workers", type=int, default=4, help="Parallel workers")
    parser.add_argument("--output-dir", default="docs/reports", help="Output directory")
    args = parser.parse_args()

    pdf_root = Path(args.root)
    pdfs = sorted(pdf_root.rglob("*.pdf"))
    if len(pdfs) < args.count:
        print(f"Not enough PDFs: {len(pdfs)}")
        return 1

    random.seed(args.seed)
    sample = random.sample(pdfs, args.count)

    tasks = [
        (pdf_path, args.dpi, args.lang, args.psm, args.page_mode, args.min_embed_tokens)
        for pdf_path in sample
    ]

    rows = []
    total = len(tasks)
    processed = 0

    if args.workers <= 1:
        for task in tasks:
            rows.append(analyze_pdf_worker(task))
            processed += 1
            if processed % 25 == 0 or processed == total:
                print(f"Processed: {processed}/{total}")
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as executor:
            futures = [executor.submit(analyze_pdf_worker, task) for task in tasks]
            for fut in as_completed(futures):
                rows.append(fut.result())
                processed += 1
                if processed % 25 == 0 or processed == total:
                    print(f"Processed: {processed}/{total}")

    summary = summarize(rows)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / f"ocr_noise_compare_{args.count}.json"
    docx_path = output_dir / f"ocr_noise_compare_{args.count}.docx"

    with json_path.open("w", encoding="utf-8") as f:
        json.dump({"summary": summary, "rows": rows}, f, ensure_ascii=False, indent=2)

    build_docx_report(summary, rows, docx_path)

    print(f"Wrote: {json_path}")
    if DOCX_AVAILABLE:
        print(f"Wrote: {docx_path}")
    else:
        print("DOCX report skipped (python-docx not installed).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
