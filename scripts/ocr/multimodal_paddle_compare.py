#!/usr/bin/env python3
"""Compare Tesseract vs PaddleOCR on mixed samples.

Runs OCR on 3 pages (first/mid/last) for each PDF and compares to embedded text.
Outputs JSON + DOCX report.

Usage:
  .venv_paddle/bin/python scripts/multimodal_paddle_compare.py \
    --candidates docs/reports/ocr_candidates_full.json \
    --candidate-count 20 --random-count 20 --seed 42 \
    --output-dir docs/reports
"""

from __future__ import annotations

import argparse
import json
import os
import random
import re
import subprocess
import tempfile
from pathlib import Path
from statistics import mean
from typing import Iterable

os.environ.setdefault("DISABLE_MODEL_SOURCE_CHECK", "True")

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from paddleocr import PaddleOCR

DEFAULT_THRESHOLDS = {
    "min_similarity": 0.25,
    "min_ocr_tokens": 40,
    "max_extracted_tokens": 20,
    "min_extracted_len": 200,
    "min_ocr_len": 200,
    "min_alpha_ratio": 0.6,
}


def load_candidates(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("candidates", [])


def sample_candidates(items: list[dict], count: int, seed: int) -> list[dict]:
    if count <= 0 or count >= len(items):
        return items
    random.seed(seed)
    return random.sample(items, count)


def list_random_pdfs(pdf_root: Path, exclude: set[str]) -> list[Path]:
    pdfs = sorted(pdf_root.rglob("*.pdf"))
    if not exclude:
        return pdfs
    return [p for p in pdfs if str(p) not in exclude]


def sample_random_pdfs(pdfs: list[Path], count: int, seed: int) -> list[Path]:
    if count <= 0 or count >= len(pdfs):
        return pdfs
    random.seed(seed)
    return random.sample(pdfs, count)


def pick_pages(page_count: int, mode: str) -> list[int]:
    if page_count <= 0:
        return []
    if mode == "first":
        return [0]
    mid = page_count // 2
    idxs = {0, mid, page_count - 1}
    return sorted(i for i in idxs if 0 <= i < page_count)


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def tokenize(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def jaccard(a: Iterable[str], b: Iterable[str]) -> float:
    set_a = set(a)
    set_b = set(b)
    if not set_a and not set_b:
        return 1.0
    if not set_a or not set_b:
        return 0.0
    return len(set_a & set_b) / len(set_a | set_b)


def alpha_ratio(text: str) -> float:
    raw = re.sub(r"\s+", "", text)
    if not raw:
        return 0.0
    good = sum(1 for ch in raw if ch.isalnum())
    return good / len(raw)


def sample_text(text: str, limit: int = 160) -> str:
    t = normalize_text(text)
    return t[:limit]


def run_tesseract(img_path: Path, lang: str, psm: int) -> str:
    cmd = [
        "tesseract",
        str(img_path),
        "stdout",
        "-l",
        lang,
        "--psm",
        str(psm),
    ]
    out = subprocess.check_output(cmd, stderr=subprocess.DEVNULL)
    return out.decode("utf-8", errors="ignore")


def ocr_page_tesseract(page: fitz.Page, dpi: int, lang: str, psm: int, temp_dir: Path) -> str:
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csGRAY)
    img_path = temp_dir / "page.png"
    pix.save(img_path)
    return run_tesseract(img_path, lang, psm)


def ocr_page_paddle(page: fitz.Page, dpi: int, ocr: PaddleOCR, temp_dir: Path) -> str:
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csGRAY)
    img_path = temp_dir / "page.png"
    pix.save(img_path)
    result = ocr.predict(str(img_path))
    if not result:
        return ""
    item = result[0]
    texts = []
    try:
        texts = item["rec_texts"]
    except Exception:
        try:
            texts = item.rec_texts
        except Exception:
            texts = []
    return "\n".join(texts)


def score_page(extracted: str, ocr_text: str, thresholds: dict) -> dict:
    extracted_tokens = tokenize(extracted)
    ocr_tokens = tokenize(ocr_text)
    sim = jaccard(extracted_tokens, ocr_tokens)
    extracted_len = len(extracted)
    ocr_len = len(ocr_text)
    alpha = alpha_ratio(extracted)
    flags = []
    if extracted_len < thresholds["min_extracted_len"] and ocr_len >= thresholds["min_ocr_len"]:
        flags.append("low_extracted_text")
    if sim < thresholds["min_similarity"] and len(ocr_tokens) >= thresholds["min_ocr_tokens"]:
        flags.append("low_similarity")
    if alpha < thresholds["min_alpha_ratio"] and extracted_len >= thresholds["min_extracted_len"]:
        flags.append("garbage_text")
    if len(ocr_tokens) >= thresholds["min_ocr_tokens"] and len(extracted_tokens) <= thresholds["max_extracted_tokens"]:
        flags.append("ocr_rich_text")

    return {
        "extracted_len": extracted_len,
        "ocr_len": ocr_len,
        "extracted_tokens": len(extracted_tokens),
        "ocr_tokens": len(ocr_tokens),
        "similarity": round(sim, 3),
        "alpha_ratio": round(alpha, 3),
        "flags": flags,
        "extracted_sample": sample_text(extracted),
        "ocr_sample": sample_text(ocr_text),
    }


def analyze_pdf(
    pdf_path: Path,
    dpi: int,
    lang: str,
    psm: int,
    thresholds: dict,
    ocr: PaddleOCR,
    page_mode: str,
    run_paddle: bool,
) -> dict:
    result = {
        "path": str(pdf_path),
        "type": pdf_path.parent.name,
        "pages": 0,
        "page_metrics": [],
        "flags": [],
        "avg_similarity_tesseract": 0.0,
        "avg_similarity_paddle": 0.0,
        "worst_page_tesseract": None,
        "worst_page_paddle": None,
        "error": None,
    }

    try:
        with fitz.open(pdf_path) as doc:
            page_count = len(doc)
            result["pages"] = page_count
            page_indexes = pick_pages(page_count, page_mode)
            if not page_indexes:
                result["error"] = "empty_pdf"
                return result

            sims_t = []
            sims_p = []

            with tempfile.TemporaryDirectory() as tmp_dir:
                temp_dir = Path(tmp_dir)
                for page_idx in page_indexes:
                    page = doc.load_page(page_idx)
                    extracted = page.get_text() or ""

                    tesseract_text = ""
                    try:
                        tesseract_text = ocr_page_tesseract(page, dpi, lang, psm, temp_dir)
                    except Exception:
                        pass

                    t_score = score_page(extracted, tesseract_text, thresholds)
                    sims_t.append(t_score["similarity"])

                    paddle_text = ""
                    p_score = None
                    if run_paddle:
                        try:
                            paddle_text = ocr_page_paddle(page, dpi, ocr, temp_dir)
                        except Exception:
                            paddle_text = ""
                        p_score = score_page(extracted, paddle_text, thresholds)
                        sims_p.append(p_score["similarity"])

                    result["page_metrics"].append({
                        "page": page_idx + 1,
                        "tesseract": t_score,
                        "paddle": p_score,
                        "paddle_skipped": not run_paddle,
                    })

            result["avg_similarity_tesseract"] = round(mean(sims_t), 3) if sims_t else 0.0
            result["avg_similarity_paddle"] = round(mean(sims_p), 3) if sims_p else None

            worst_t = min(result["page_metrics"], key=lambda m: m["tesseract"]["similarity"])
            if sims_p:
                worst_p = min(
                    [m for m in result["page_metrics"] if m["paddle"] is not None],
                    key=lambda m: m["paddle"]["similarity"],
                )
            else:
                worst_p = None
            result["worst_page_tesseract"] = worst_t
            result["worst_page_paddle"] = worst_p

    except Exception as exc:
        result["error"] = str(exc)

    return result


def compute_stats(rows: list[dict]) -> dict:
    if not rows:
        return {"count": 0, "avg_similarity_tesseract": 0.0, "avg_similarity_paddle": 0.0}
    paddle_vals = [r["avg_similarity_paddle"] for r in rows if r["avg_similarity_paddle"] is not None]
    return {
        "count": len(rows),
        "avg_similarity_tesseract": round(mean(r["avg_similarity_tesseract"] for r in rows), 3),
        "avg_similarity_paddle": round(mean(paddle_vals), 3) if paddle_vals else None,
    }


def build_paddle_ocr(lang: str, disable_doc_preprocess: bool) -> tuple[PaddleOCR, bool]:
    if not disable_doc_preprocess:
        return PaddleOCR(lang=lang), False
    try:
        return (
            PaddleOCR(
                lang=lang,
                use_doc_orientation_classify=False,
                use_doc_unwarping=False,
            ),
            True,
        )
    except TypeError:
        try:
            return (PaddleOCR(lang=lang, use_doc_orientation_classify=False), True)
        except TypeError:
            return PaddleOCR(lang=lang), False


def build_docx_report(summary: dict, rows: list[dict], output_path: Path) -> None:
    doc = Document()
    total = summary.get("files_total", 0)
    title = doc.add_heading(f"Tesseract vs PaddleOCR 비교 보고서 (혼합 샘플 {total}건)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cand = summary.get("candidate_count", 0)
    rand = summary.get("random_count", 0)
    doc.add_paragraph(
        f"설명: OCR 후보군 {cand}건 + 일반 무작위 {rand}건(총 {total}건)을 혼합 분석했습니다."
    )
    page_mode = summary.get("page_mode", "three")
    if page_mode == "first":
        doc.add_paragraph("방법: 각 문서의 1페이지를 OCR로 읽고,")
    else:
        doc.add_paragraph("방법: 각 문서의 1/중간/마지막 페이지를 OCR로 읽고,")
    doc.add_paragraph("      같은 페이지의 PDF 내장 텍스트와 유사도를 계산했습니다.")
    doc.add_paragraph(
        f"설정: dpi={summary.get('dpi')} | paddle_groups={summary.get('paddle_groups')} | "
        f"doc_preprocess_disabled={summary.get('doc_preprocess_disabled')}"
    )

    doc.add_heading("1. 핵심 요약", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "값"
    for label, value in [
        ("총 샘플 수", summary.get("files_total")),
        ("분석 성공", summary.get("files_ok")),
        ("분석 오류", summary.get("files_error")),
    ]:
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(value)

    doc.add_heading("2. Tesseract vs PaddleOCR (평균 유사도)", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "그룹"
    table.rows[0].cells[1].text = "Tesseract"
    table.rows[0].cells[2].text = "PaddleOCR"
    for group_key, label in (("candidate", "후보군"), ("random", "일반군"), ("overall", "전체")):
        stats = summary.get("group_stats", {}).get(group_key, {})
        paddle_val = stats.get("avg_similarity_paddle")
        paddle_text = "n/a" if paddle_val is None else str(paddle_val)
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(stats.get("avg_similarity_tesseract", 0.0))
        row[2].text = paddle_text

    doc.add_heading("3. 의심 사례 (유사도 최저 5건)", level=1)
    for row in summary.get("worst_cases", []):
        doc.add_paragraph(f"{row['group']} | {row['path']}", style="List Bullet")
        doc.add_paragraph(
            f"  tesseract_avg={row['avg_similarity_tesseract']} | paddle_avg={row['avg_similarity_paddle']}"
        )
        doc.add_paragraph(
            f"  tesseract_sample: {row['tesseract_sample']}"
        )
        doc.add_paragraph(
            f"  paddle_sample: {row['paddle_sample']}"
        )

    doc.add_heading("4. 기준값(임계치)", level=1)
    for k, v in summary.get("thresholds", {}).items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="PaddleOCR vs Tesseract comparison")
    parser.add_argument("--candidates", required=True, help="OCR candidates JSON")
    parser.add_argument("--candidate-count", type=int, default=20, help="Candidate sample size")
    parser.add_argument("--random-root", default="peraturan/data/pdfs", help="PDF root for random sample")
    parser.add_argument("--random-count", type=int, default=20, help="Random sample size")
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument("--dpi", type=int, default=200, help="Render DPI for OCR")
    parser.add_argument("--lang", default="id", help="PaddleOCR language")
    parser.add_argument("--tess-lang", default="ind+eng", help="Tesseract languages")
    parser.add_argument("--psm", type=int, default=6, help="Tesseract PSM")
    parser.add_argument("--page-mode", choices=["first", "three"], default="three", help="Pages to OCR")
    parser.add_argument(
        "--paddle-groups",
        choices=["all", "candidate", "none"],
        default="all",
        help="Which groups use PaddleOCR",
    )
    parser.add_argument("--disable-doc-preprocess", action="store_true", help="Disable doc preprocessor")
    parser.add_argument("--tag", default="", help="Tag to append to output files")
    parser.add_argument("--output-dir", default="docs/reports", help="Output directory")
    args = parser.parse_args()

    candidates_path = Path(args.candidates)
    candidate_items = load_candidates(candidates_path)
    if not candidate_items:
        print("No candidates found.")
        return 1

    candidate_sample = sample_candidates(candidate_items, args.candidate_count, args.seed)
    candidate_set = {item["path"] for item in candidate_items}

    pdf_root = Path(args.random_root)
    random_pool = list_random_pdfs(pdf_root, exclude=candidate_set)
    if len(random_pool) < args.random_count:
        random_pool = list_random_pdfs(pdf_root, exclude=set())

    random_sample = sample_random_pdfs(random_pool, args.random_count, args.seed + 1)

    ocr, doc_pre_disabled = build_paddle_ocr(args.lang, args.disable_doc_preprocess)

    rows = []
    total = len(candidate_sample) + len(random_sample)
    processed = 0
    for item in candidate_sample:
        run_paddle = args.paddle_groups in ("all", "candidate")
        rows.append(
            analyze_pdf(
                Path(item["path"]),
                args.dpi,
                args.tess_lang,
                args.psm,
                DEFAULT_THRESHOLDS,
                ocr,
                args.page_mode,
                run_paddle,
            )
            | {"group": "candidate"}
        )
        processed += 1
        if processed % 5 == 0 or processed == total:
            print(f"Processed {processed}/{total}")
    for pdf_path in random_sample:
        run_paddle = args.paddle_groups == "all"
        rows.append(
            analyze_pdf(
                Path(pdf_path),
                args.dpi,
                args.tess_lang,
                args.psm,
                DEFAULT_THRESHOLDS,
                ocr,
                args.page_mode,
                run_paddle,
            )
            | {"group": "random"}
        )
        processed += 1
        if processed % 5 == 0 or processed == total:
            print(f"Processed {processed}/{total}")

    ok_rows = [r for r in rows if not r.get("error")]
    error_rows = [r for r in rows if r.get("error")]
    candidate_rows = [r for r in ok_rows if r.get("group") == "candidate"]
    random_rows = [r for r in ok_rows if r.get("group") == "random"]

    summary = {
        "files_total": len(rows),
        "files_ok": len(ok_rows),
        "files_error": len(error_rows),
        "candidate_count": len(candidate_rows),
        "random_count": len(random_rows),
        "page_mode": args.page_mode,
        "paddle_groups": args.paddle_groups,
        "doc_preprocess_disabled": doc_pre_disabled,
        "dpi": args.dpi,
        "thresholds": DEFAULT_THRESHOLDS,
    }

    summary["group_stats"] = {
        "candidate": compute_stats(candidate_rows),
        "random": compute_stats(random_rows),
        "overall": compute_stats(ok_rows),
    }

    def min_similarity(row: dict) -> float:
        vals = [row.get("avg_similarity_tesseract", 1.0)]
        if row.get("avg_similarity_paddle") is not None:
            vals.append(row.get("avg_similarity_paddle"))
        return min(vals)

    worst_cases = []
    for row in sorted(ok_rows, key=min_similarity)[:5]:
        wp_t = row.get("worst_page_tesseract") or {}
        wp_p = row.get("worst_page_paddle") or {}
        worst_cases.append({
            "group": row.get("group"),
            "path": row["path"],
            "avg_similarity_tesseract": row.get("avg_similarity_tesseract"),
            "avg_similarity_paddle": row.get("avg_similarity_paddle"),
            "tesseract_sample": (wp_t.get("tesseract", {}) or {}).get("ocr_sample"),
            "paddle_sample": (wp_p.get("paddle", {}) or {}).get("ocr_sample") if wp_p else None,
        })
    summary["worst_cases"] = worst_cases

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    total_count = len(rows)
    suffix = "p1" if args.page_mode == "first" else "p3"
    tag = f"_{args.tag}" if args.tag else ""
    json_path = output_dir / f"paddle_vs_tesseract_{total_count}_{suffix}{tag}.json"
    docx_path = output_dir / f"paddle_vs_tesseract_{total_count}_{suffix}{tag}.docx"

    with json_path.open("w", encoding="utf-8") as f:
        json.dump({"summary": summary, "rows": rows}, f, ensure_ascii=False, indent=2)

    build_docx_report(summary, rows, docx_path)

    print(f"Wrote: {json_path}")
    print(f"Wrote: {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
