#!/usr/bin/env python3
"""Mixed sample report: OCR candidates + random PDFs.

Runs OCR-vs-embedded-text comparison on 3 pages (first/middle/last),
for a mixed sample (e.g., 500 high-risk + 500 random).

Usage:
  .venv/bin/python scripts/multimodal_mixed_report.py \
    --candidates docs/reports/ocr_candidates_full.json \
    --candidate-count 500 --random-count 500 --seed 42 \
    --output-dir docs/reports
"""

from __future__ import annotations

import argparse
import json
import random
import re
import subprocess
import tempfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path
from statistics import mean
from typing import Iterable

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


DEFAULT_THRESHOLDS = {
    "min_similarity": 0.25,
    "min_ocr_tokens": 40,
    "max_extracted_tokens": 20,
    "min_extracted_len": 200,
    "min_ocr_len": 200,
    "min_alpha_ratio": 0.6,
}


def load_candidates(path: Path) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("candidates", [])


def sample_candidates(items: list[dict], count: int, seed: int) -> list[dict]:
    if count <= 0:
        return []
    if count >= len(items):
        return items
    random.seed(seed)
    return random.sample(items, count)


def list_random_pdfs(pdf_root: Path, exclude: set[str]) -> list[Path]:
    pdfs = sorted(pdf_root.rglob("*.pdf"))
    if not exclude:
        return pdfs
    return [p for p in pdfs if str(p) not in exclude]


def sample_random_pdfs(pdfs: list[Path], count: int, seed: int) -> list[Path]:
    if count <= 0 or count >= len(pdfs):
        return pdfs
    random.seed(seed)
    return random.sample(pdfs, count)


def pick_pages(page_count: int) -> list[int]:
    if page_count <= 0:
        return []
    mid = page_count // 2
    idxs = {0, mid, page_count - 1}
    return sorted(i for i in idxs if 0 <= i < page_count)


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def tokenize(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


def jaccard(a: Iterable[str], b: Iterable[str]) -> float:
    set_a = set(a)
    set_b = set(b)
    if not set_a and not set_b:
        return 1.0
    if not set_a or not set_b:
        return 0.0
    return len(set_a & set_b) / len(set_a | set_b)


def alpha_ratio(text: str) -> float:
    raw = re.sub(r"\s+", "", text)
    if not raw:
        return 0.0
    good = sum(1 for ch in raw if ch.isalnum())
    return good / len(raw)


def sample_text(text: str, limit: int = 180) -> str:
    t = normalize_text(text)
    return t[:limit]


def run_tesseract(img_path: Path, lang: str, psm: int) -> str:
    cmd = [
        "tesseract",
        str(img_path),
        "stdout",
        "-l",
        lang,
        "--psm",
        str(psm),
    ]
    out = subprocess.check_output(cmd, stderr=subprocess.DEVNULL)
    return out.decode("utf-8", errors="ignore")


def ocr_page(page: fitz.Page, dpi: int, lang: str, psm: int, temp_dir: Path) -> str:
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False, colorspace=fitz.csGRAY)
    img_path = temp_dir / "page.png"
    pix.save(img_path)
    return run_tesseract(img_path, lang, psm)


def analyze_pdf(
    pdf_path: Path,
    dpi: int,
    lang: str,
    psm: int,
    thresholds: dict,
) -> dict:
    result = {
        "path": str(pdf_path),
        "type": pdf_path.parent.name,
        "pages": 0,
        "page_metrics": [],
        "flags": [],
        "avg_similarity": 0.0,
        "worst_page": None,
        "error": None,
    }

    try:
        with fitz.open(pdf_path) as doc:
            page_count = len(doc)
            result["pages"] = page_count
            page_indexes = pick_pages(page_count)
            if not page_indexes:
                result["error"] = "empty_pdf"
                return result

            similarities = []
            low_similarity_pages = 0
            low_text_pages = 0
            garbage_pages = 0
            ocr_rich_pages = 0

            with tempfile.TemporaryDirectory() as tmp_dir:
                temp_dir = Path(tmp_dir)
                for page_idx in page_indexes:
                    page = doc.load_page(page_idx)
                    extracted = page.get_text() or ""
                    ocr_text = ""
                    ocr_error = None
                    try:
                        ocr_text = ocr_page(page, dpi, lang, psm, temp_dir)
                    except Exception as exc:
                        ocr_error = str(exc)

                    extracted_tokens = tokenize(extracted)
                    ocr_tokens = tokenize(ocr_text)
                    sim = jaccard(extracted_tokens, ocr_tokens)
                    similarities.append(sim)

                    extracted_len = len(extracted)
                    ocr_len = len(ocr_text)
                    alpha = alpha_ratio(extracted)

                    flags = []
                    if ocr_error:
                        flags.append("ocr_error")
                    if extracted_len < thresholds["min_extracted_len"] and ocr_len >= thresholds["min_ocr_len"]:
                        flags.append("low_extracted_text")
                        low_text_pages += 1
                    if sim < thresholds["min_similarity"] and len(ocr_tokens) >= thresholds["min_ocr_tokens"]:
                        flags.append("low_similarity")
                        low_similarity_pages += 1
                    if alpha < thresholds["min_alpha_ratio"] and extracted_len >= thresholds["min_extracted_len"]:
                        flags.append("garbage_text")
                        garbage_pages += 1
                    if len(ocr_tokens) >= thresholds["min_ocr_tokens"] and len(extracted_tokens) <= thresholds["max_extracted_tokens"]:
                        flags.append("ocr_rich_text")
                        ocr_rich_pages += 1

                    metric = {
                        "page": page_idx + 1,
                        "extracted_len": extracted_len,
                        "ocr_len": ocr_len,
                        "extracted_tokens": len(extracted_tokens),
                        "ocr_tokens": len(ocr_tokens),
                        "similarity": round(sim, 3),
                        "alpha_ratio": round(alpha, 3),
                        "flags": flags,
                        "extracted_sample": sample_text(extracted),
                        "ocr_sample": sample_text(ocr_text),
                    }
                    result["page_metrics"].append(metric)

            result["avg_similarity"] = round(mean(similarities), 3) if similarities else 0.0
            worst = min(result["page_metrics"], key=lambda m: m["similarity"]) if result["page_metrics"] else None
            result["worst_page"] = worst

            if low_similarity_pages > 0:
                result["flags"].append("low_similarity_doc")
            if low_text_pages > 0:
                result["flags"].append("low_extracted_text_doc")
            if garbage_pages > 0:
                result["flags"].append("garbage_text_doc")
            if ocr_rich_pages > 0:
                result["flags"].append("ocr_rich_text_doc")

    except Exception as exc:
        result["error"] = str(exc)

    return result


def analyze_pdf_worker(args: tuple[Path, str, int, str, int, dict]) -> dict:
    pdf_path, group, dpi, lang, psm, thresholds = args
    result = analyze_pdf(pdf_path, dpi, lang, psm, thresholds)
    result["group"] = group
    return result


def compute_stats(rows: list[dict]) -> dict:
    if not rows:
        return {
            "count": 0,
            "avg_similarity": 0.0,
            "low_similarity": 0,
            "ocr_rich": 0,
            "garbage_text": 0,
        }
    return {
        "count": len(rows),
        "avg_similarity": round(mean(r["avg_similarity"] for r in rows), 3),
        "low_similarity": sum(1 for r in rows if "low_similarity_doc" in r.get("flags", [])),
        "ocr_rich": sum(1 for r in rows if "ocr_rich_text_doc" in r.get("flags", [])),
        "garbage_text": sum(1 for r in rows if "garbage_text_doc" in r.get("flags", [])),
    }


def build_docx_report(summary: dict, rows: list[dict], output_path: Path) -> None:
    doc = Document()
    total = summary.get("files_total", 0)
    candidate_count = summary.get("candidate_count", 0)
    random_count = summary.get("random_count", 0)
    if candidate_count == 0:
        title = doc.add_heading(f"무작위 샘플 {total}건 이미지-텍스트 비교 보고서", level=0)
    else:
        title = doc.add_heading(f"혼합 샘플 {total}건 이미지-텍스트 비교 보고서", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if candidate_count == 0:
        doc.add_paragraph(f"설명: 전체 PDF에서 무작위 {random_count}건을 추출해 분석했습니다.")
    else:
        doc.add_paragraph(
            f"설명: OCR 후보군 {candidate_count}건 + 일반 무작위 {random_count}건을 혼합 분석했습니다."
        )
    doc.add_paragraph("방법: 각 문서의 1/중간/마지막 페이지를 OCR로 읽고,")
    doc.add_paragraph("      같은 페이지의 PDF 내장 텍스트와 유사도를 계산했습니다.")

    doc.add_heading("1. 핵심 요약", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "항목"
    table.rows[0].cells[1].text = "값"
    summary_rows = [
        ("총 샘플 수", summary.get("files_total")),
        ("분석 성공", summary.get("files_ok")),
        ("분석 오류", summary.get("files_error")),
    ]
    if candidate_count > 0:
        summary_rows.insert(1, ("후보군 샘플 수", summary.get("candidate_count")))
        summary_rows.insert(2, ("일반군 샘플 수", summary.get("random_count")))
    for label, value in summary_rows:
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(value)

    if candidate_count == 0:
        doc.add_heading("2. 전체 요약", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "그룹"
        table.rows[0].cells[1].text = "평균 유사도"
        table.rows[0].cells[2].text = "유사도 낮음"
        table.rows[0].cells[3].text = "OCR 풍부/추출 빈약"
        stats = summary.get("group_stats", {}).get("overall", {})
        row = table.add_row().cells
        row[0].text = "전체"
        row[1].text = str(stats.get("avg_similarity", 0.0))
        row[2].text = str(stats.get("low_similarity", 0))
        row[3].text = str(stats.get("ocr_rich", 0))
    else:
        doc.add_heading("2. 후보군 vs 일반군 비교", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "그룹"
        table.rows[0].cells[1].text = "평균 유사도"
        table.rows[0].cells[2].text = "유사도 낮음"
        table.rows[0].cells[3].text = "OCR 풍부/추출 빈약"
        for group_key, label in (("candidate", "후보군"), ("random", "일반군"), ("overall", "전체")):
            stats = summary.get("group_stats", {}).get(group_key, {})
            row = table.add_row().cells
            row[0].text = label
            row[1].text = str(stats.get("avg_similarity", 0.0))
            row[2].text = str(stats.get("low_similarity", 0))
            row[3].text = str(stats.get("ocr_rich", 0))

    doc.add_heading("3. 해석", level=1)
    for line in summary.get("interpretation", []):
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("4. 유형별 분포 (참고)", level=1)
    type_breakdown = summary.get("type_breakdown", {})
    if type_breakdown:
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = "유형"
        table.rows[0].cells[1].text = "건수"
        table.rows[0].cells[2].text = "유사도 낮음"
        table.rows[0].cells[3].text = "OCR 풍부/추출 빈약"
        for key, stats in type_breakdown.items():
            row = table.add_row().cells
            row[0].text = key
            row[1].text = str(stats.get("count", 0))
            row[2].text = str(stats.get("low_similarity", 0))
            row[3].text = str(stats.get("ocr_rich", 0))
    else:
        doc.add_paragraph("유형별 통계가 없습니다.")

    doc.add_heading("5. 의심 사례 (유사도 최저 10건)", level=1)
    for row in summary.get("worst_cases", []):
        doc.add_paragraph(f"{row['group']} | {row['path']}", style="List Bullet")
        doc.add_paragraph(
            f"  avg_similarity={row['avg_similarity']} | flags={', '.join(row['flags'])}"
        )
        doc.add_paragraph(
            f"  worst_page={row['page']} | sim={row['similarity']} | "
            f"extracted_len={row['extracted_len']} | ocr_len={row['ocr_len']}"
        )
        doc.add_paragraph(f"  extracted_sample: {row['extracted_sample']}")
        doc.add_paragraph(f"  ocr_sample: {row['ocr_sample']}")

    doc.add_heading("6. 기준값(임계치)", level=1)
    thresholds = summary.get("thresholds", {})
    for k, v in thresholds.items():
        doc.add_paragraph(f"{k}: {v}", style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Mixed sample multimodal report")
    parser.add_argument("--candidates", required=True, help="OCR candidates JSON")
    parser.add_argument("--candidate-count", type=int, default=500, help="Candidate sample size")
    parser.add_argument("--random-root", default="peraturan/data/pdfs", help="PDF root for random sample")
    parser.add_argument("--random-count", type=int, default=500, help="Random sample size")
    parser.add_argument(
        "--random-include-candidates",
        action="store_true",
        help="Include candidate PDFs in random pool",
    )
    parser.add_argument("--seed", type=int, default=42, help="Random seed")
    parser.add_argument("--dpi", type=int, default=200, help="Render DPI for OCR")
    parser.add_argument("--lang", default="ind+eng", help="Tesseract languages")
    parser.add_argument("--psm", type=int, default=6, help="Tesseract PSM")
    parser.add_argument("--workers", type=int, default=4, help="Parallel workers")
    parser.add_argument("--output-dir", default="docs/reports", help="Output directory")
    args = parser.parse_args()

    candidates_path = Path(args.candidates)
    candidate_items = load_candidates(candidates_path)
    if not candidate_items:
        print("No candidates found.")
        return 1

    candidate_sample = sample_candidates(candidate_items, args.candidate_count, args.seed)
    candidate_set = set() if args.random_include_candidates else {item["path"] for item in candidate_items}

    pdf_root = Path(args.random_root)
    random_pool = list_random_pdfs(pdf_root, exclude=candidate_set)
    overlap_used = False
    if len(random_pool) < args.random_count:
        random_pool = list_random_pdfs(pdf_root, exclude=set())
        overlap_used = True

    random_sample = sample_random_pdfs(random_pool, args.random_count, args.seed + 1)

    tasks = []
    for item in candidate_sample:
        tasks.append((Path(item["path"]), "candidate", args.dpi, args.lang, args.psm, DEFAULT_THRESHOLDS))
    for pdf_path in random_sample:
        tasks.append((Path(pdf_path), "random", args.dpi, args.lang, args.psm, DEFAULT_THRESHOLDS))

    rows = []
    total = len(tasks)
    processed = 0

    if args.workers <= 1:
        for task in tasks:
            rows.append(analyze_pdf_worker(task))
            processed += 1
            if processed % 50 == 0:
                print(f"Processed: {processed}/{total}")
    else:
        with ProcessPoolExecutor(max_workers=args.workers) as executor:
            futures = [executor.submit(analyze_pdf_worker, task) for task in tasks]
            for fut in as_completed(futures):
                rows.append(fut.result())
                processed += 1
                if processed % 50 == 0:
                    print(f"Processed: {processed}/{total}")

    ok_rows = [r for r in rows if not r.get("error")]
    error_rows = [r for r in rows if r.get("error")]
    candidate_rows = [r for r in ok_rows if r.get("group") == "candidate"]
    random_rows = [r for r in ok_rows if r.get("group") == "random"]

    summary = {
        "files_total": len(rows),
        "files_ok": len(ok_rows),
        "files_error": len(error_rows),
        "candidate_count": len(candidate_rows),
        "random_count": len(random_rows),
        "overlap_used": overlap_used,
        "random_include_candidates": args.random_include_candidates,
        "thresholds": DEFAULT_THRESHOLDS,
    }

    summary["group_stats"] = {
        "candidate": compute_stats(candidate_rows),
        "random": compute_stats(random_rows),
        "overall": compute_stats(ok_rows),
    }

    interpretation = []
    cand = summary["group_stats"]["candidate"]
    rand = summary["group_stats"]["random"]

    if cand["low_similarity"] > 0:
        interpretation.append("후보군에서 OCR-텍스트 불일치가 뚜렷하게 나타났습니다.")
    if rand["low_similarity"] > 0:
        interpretation.append("일반군에서도 불일치 문서가 존재합니다.")
    if rand["low_similarity"] > (0.2 * max(1, rand["count"])):
        interpretation.append("일반군에서도 OCR 필요 문서 비율이 높아 보입니다. 범위 확대 검토가 필요합니다.")
    if rand["low_similarity"] == 0:
        interpretation.append("일반군은 대부분 내장 텍스트 활용이 가능해 보입니다.")
    if overlap_used:
        interpretation.append("일반군 표본이 부족하여 일부 중복 샘플이 포함되었습니다.")
    summary["interpretation"] = interpretation

    type_breakdown = {}
    by_type = {}
    for row in ok_rows:
        by_type.setdefault(row["type"], []).append(row)
    for key, items in sorted(by_type.items(), key=lambda x: len(x[1]), reverse=True):
        type_breakdown[key] = {
            "count": len(items),
            "low_similarity": sum(1 for r in items if "low_similarity_doc" in r.get("flags", [])),
            "ocr_rich": sum(1 for r in items if "ocr_rich_text_doc" in r.get("flags", [])),
        }
    summary["type_breakdown"] = type_breakdown

    worst_cases = []
    for row in sorted(ok_rows, key=lambda r: r.get("avg_similarity", 1.0))[:10]:
        wp = row.get("worst_page") or {}
        worst_cases.append({
            "group": row.get("group"),
            "path": row["path"],
            "avg_similarity": row.get("avg_similarity"),
            "flags": row.get("flags", []),
            "page": wp.get("page"),
            "similarity": wp.get("similarity"),
            "extracted_len": wp.get("extracted_len"),
            "ocr_len": wp.get("ocr_len"),
            "extracted_sample": wp.get("extracted_sample"),
            "ocr_sample": wp.get("ocr_sample"),
        })
    summary["worst_cases"] = worst_cases

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    total = len(rows)
    if summary.get("candidate_count", 0) == 0:
        json_path = output_dir / f"multimodal_random_{total}.json"
        docx_path = output_dir / f"multimodal_random_{total}.docx"
    else:
        json_path = output_dir / f"multimodal_mixed_{total}.json"
        docx_path = output_dir / f"multimodal_mixed_{total}.docx"

    with json_path.open("w", encoding="utf-8") as f:
        json.dump({"summary": summary, "rows": rows}, f, ensure_ascii=False, indent=2)

    build_docx_report(summary, rows, docx_path)

    print(f"Wrote: {json_path}")
    print(f"Wrote: {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
