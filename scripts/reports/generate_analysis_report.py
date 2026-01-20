#!/usr/bin/env python3
"""
인도네시아 법령정보시스템 구축 기초자료 - 종합분석 보고서 생성 스크립트

대상: peraturan.go.id (법제처) - RFP 대상 법령 6개 유형만
"""

import json
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, RGBColor

# === 경로 설정 ===
PROJECT_ROOT = Path(__file__).parent.parent.parent
PERATURAN_DB = PROJECT_ROOT / "peraturan" / "data" / "peraturan.db"
REFERENCES_JSON = PROJECT_ROOT / "peraturan" / "data" / "all_references.json"
OUTPUT_DIR = PROJECT_ROOT / "docs" / "exports"
OUTPUT_FILE = OUTPUT_DIR / "법령정보시스템_기초자료_분석보고서.docx"

# === RFP 대상 법령 유형 (6개) ===
RFP_TYPE_ABBREV = {
    "UNDANG-UNDANG": "UU",
    "PERATURAN PEMERINTAH": "PP",
    "PERATURAN PRESIDEN": "PERPRES",
    "PERATURAN PEMERINTAH PENGGANTI UNDANG-UNDANG": "PERPPU",
    "PERATURAN MENTERI": "PERMEN",
    "PERATURAN BADAN/LEMBAGA": "PERBAN",
}

RFP_TYPE_KOREAN = {
    "UNDANG-UNDANG": "법률",
    "PERATURAN PEMERINTAH": "정부령",
    "PERATURAN PRESIDEN": "대통령령",
    "PERATURAN PEMERINTAH PENGGANTI UNDANG-UNDANG": "긴급법률",
    "PERATURAN MENTERI": "장관령",
    "PERATURAN BADAN/LEMBAGA": "기관규정",
}

RFP_WHERE = "jenis IN ('UNDANG-UNDANG','PERATURAN PEMERINTAH','PERATURAN PRESIDEN','PERATURAN PEMERINTAH PENGGANTI UNDANG-UNDANG','PERATURAN MENTERI','PERATURAN BADAN/LEMBAGA')"


def query_db(db_path: Path, query: str) -> list[tuple]:
    if not db_path.exists():
        return []
    conn = sqlite3.connect(db_path)
    cursor = conn.execute(query)
    results = cursor.fetchall()
    conn.close()
    return results


def query_one(db_path: Path, query: str) -> Any:
    result = query_db(db_path, query)
    return result[0][0] if result else 0


def fmt(n: int | float) -> str:
    if n is None:
        return "0"
    if isinstance(n, float):
        return f"{n:,.1f}"
    return f"{n:,}"


def pct(part: int, total: int) -> str:
    if total == 0:
        return "0%"
    return f"{part / total * 100:.1f}%"


def set_cell_shading(cell, color: str):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def create_table(doc: Document, data: list[list[str]], header: bool = True) -> Any:
    if not data:
        return None
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = str(cell_text)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if header and i == 0:
                set_cell_shading(cell, "4472C4")
                for run in para.runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    return table


class ReportGenerator:
    def __init__(self):
        self.doc = Document()
        self.stats = self._load_stats()
        self._setup_styles()

    def _setup_styles(self):
        styles = self.doc.styles
        for style_name in ["Title", "Normal", "Heading 1", "Heading 2"]:
            styles[style_name].font.name = "맑은 고딕"
        styles["Title"].font.size = Pt(24)
        styles["Normal"].font.size = Pt(11)
        styles["Heading 1"].font.size = Pt(16)
        styles["Heading 1"].font.color.rgb = RGBColor(0, 51, 102)
        styles["Heading 2"].font.size = Pt(14)

    def _load_stats(self) -> dict[str, Any]:
        s = {}

        # 기본 통계
        s["total"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE}")
        s["year_min"], s["year_max"] = query_db(PERATURAN_DB, f"SELECT MIN(tahun), MAX(tahun) FROM peraturan WHERE {RFP_WHERE}")[0]

        # PDF 통계
        s["pdf_url_exists"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND pdf_url IS NOT NULL")
        s["pdf_url_none"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND pdf_url IS NULL")
        s["pdf_downloaded"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND local_pdf_path IS NOT NULL")
        s["pdf_download_failed"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND pdf_url IS NOT NULL AND local_pdf_path IS NULL")

        # 텍스트 추출 통계
        s["extract_success"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND extraction_success = 1")
        s["extract_needs_ocr"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND needs_ocr = 1")
        s["extract_failed"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND local_pdf_path IS NOT NULL AND extraction_success = 0 AND needs_ocr = 0")

        # 파싱 통계
        s["parse_success"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND parse_success = 1")
        s["total_bab"] = query_one(PERATURAN_DB, f"SELECT SUM(parsed_bab_count) FROM peraturan WHERE {RFP_WHERE}") or 0
        s["total_pasal"] = query_one(PERATURAN_DB, f"SELECT SUM(parsed_pasal_count) FROM peraturan WHERE {RFP_WHERE}") or 0
        s["total_ayat"] = query_one(PERATURAN_DB, f"SELECT SUM(parsed_ayat_count) FROM peraturan WHERE {RFP_WHERE}") or 0
        s["total_huruf"] = query_one(PERATURAN_DB, f"SELECT SUM(parsed_huruf_count) FROM peraturan WHERE {RFP_WHERE}") or 0
        s["total_chars"] = query_one(PERATURAN_DB, f"SELECT SUM(LENGTH(extracted_text)) FROM peraturan WHERE {RFP_WHERE} AND extracted_text IS NOT NULL") or 0

        # 유형별
        s["by_type"] = query_db(PERATURAN_DB, f"SELECT jenis, COUNT(*) FROM peraturan WHERE {RFP_WHERE} GROUP BY jenis ORDER BY COUNT(*) DESC")
        s["status_berlaku"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND status = 'Berlaku'")
        s["status_tidak"] = query_one(PERATURAN_DB, f"SELECT COUNT(*) FROM peraturan WHERE {RFP_WHERE} AND status LIKE 'Tidak%'")

        s["type_pdf"] = query_db(PERATURAN_DB, f"""
            SELECT jenis, COUNT(*),
                   SUM(CASE WHEN pdf_url IS NOT NULL THEN 1 ELSE 0 END),
                   SUM(CASE WHEN local_pdf_path IS NOT NULL THEN 1 ELSE 0 END),
                   SUM(CASE WHEN extraction_success = 1 THEN 1 ELSE 0 END)
            FROM peraturan WHERE {RFP_WHERE} GROUP BY jenis ORDER BY COUNT(*) DESC
        """)

        # 연도별 생산량
        s["by_year"] = query_db(PERATURAN_DB, f"SELECT tahun, COUNT(*) FROM peraturan WHERE {RFP_WHERE} GROUP BY tahun ORDER BY COUNT(*) DESC LIMIT 10")

        # 발급기관별
        s["by_agency"] = query_db(PERATURAN_DB, f"""
            SELECT pemrakarsa, COUNT(*) FROM peraturan
            WHERE {RFP_WHERE} AND pemrakarsa IS NOT NULL AND pemrakarsa != ''
            GROUP BY pemrakarsa ORDER BY COUNT(*) DESC LIMIT 10
        """)

        # 유형별 평균 조항 수
        s["avg_pasal_by_type"] = query_db(PERATURAN_DB, f"""
            SELECT jenis, COUNT(*), ROUND(AVG(parsed_pasal_count), 1), ROUND(AVG(parsed_ayat_count), 1), MAX(parsed_pasal_count)
            FROM peraturan WHERE {RFP_WHERE} AND parse_success = 1
            GROUP BY jenis ORDER BY AVG(parsed_pasal_count) DESC
        """)

        # 문서 길이 분포
        s["length_dist"] = query_db(PERATURAN_DB, f"""
            SELECT CASE
                WHEN LENGTH(extracted_text) < 5000 THEN '5천자 미만'
                WHEN LENGTH(extracted_text) < 20000 THEN '5천~2만자'
                WHEN LENGTH(extracted_text) < 50000 THEN '2만~5만자'
                WHEN LENGTH(extracted_text) < 100000 THEN '5만~10만자'
                ELSE '10만자 이상'
            END, COUNT(*)
            FROM peraturan WHERE {RFP_WHERE} AND extracted_text IS NOT NULL
            GROUP BY 1 ORDER BY LENGTH(extracted_text)
        """)

        # 가장 긴 법령
        s["longest_docs"] = query_db(PERATURAN_DB, f"""
            SELECT jenis, nomor, tahun, substr(tentang, 1, 50), parsed_pasal_count
            FROM peraturan WHERE {RFP_WHERE} AND parse_success = 1
            ORDER BY parsed_pasal_count DESC LIMIT 5
        """)

        # 참조 분석
        s["references"] = self._load_reference_stats()

        return s

    def _load_reference_stats(self) -> dict:
        if not REFERENCES_JSON.exists():
            return {}

        with open(REFERENCES_JSON, 'r', encoding='utf-8') as f:
            data = json.load(f)

        total_ext = 0
        total_int = 0
        total_cond = 0
        ref_types = Counter()
        most_cited = Counter()

        for item in data:
            ext = item.get('external_refs', [])
            int_refs = item.get('internal_refs', [])
            cond = item.get('conditional_clauses', [])

            total_ext += len(ext)
            total_int += len(int_refs)
            total_cond += len(cond)

            for ref in ext:
                ref_types[ref.get('ref_type', 'UNKNOWN')] += 1
                cited = f"{ref.get('jenis', '')} No.{ref.get('nomor', '')} Tahun {ref.get('tahun', '')}"
                most_cited[cited] += 1

        return {
            "total_docs": len(data),
            "total_external": total_ext,
            "total_internal": total_int,
            "total_conditional": total_cond,
            "ref_types": ref_types.most_common(5),
            "most_cited": most_cited.most_common(10),
        }

    def add_cover(self):
        for _ in range(4):
            self.doc.add_paragraph()

        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("인도네시아 법령정보시스템 구축 기초자료")
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)

        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run("종합분석 보고서")
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(0, 102, 153)

        self.doc.add_paragraph()
        line = self.doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.add_run("━" * 40).font.color.rgb = RGBColor(150, 150, 150)

        self.doc.add_paragraph()
        info = self.doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info.add_run(f"작성일: {datetime.now().strftime('%Y년 %m월 %d일')}\n").font.size = Pt(12)
        info.add_run("데이터 출처: peraturan.go.id\n").font.size = Pt(12)
        info.add_run("분석 대상: RFP 대상 법령 6개 유형").font.size = Pt(12)

        self.doc.add_page_break()

    def add_toc(self):
        self.doc.add_heading("목차", level=1)
        toc = [
            "1. 요약 (Executive Summary)",
            "2. 데이터 수집 현황",
            "3. PDF 파일 현황",
            "4. 텍스트 추출 현황",
            "5. 문서 구조 파싱",
            "6. 참조/인용 분석",
            "7. 발급기관 및 연도별 분석",
            "8. 현행성 판단",
            "9. 문제점 및 권고사항",
            "부록: 용어 대조표",
        ]
        for item in toc:
            p = self.doc.add_paragraph()
            p.add_run(item).bold = True
        self.doc.add_page_break()

    def add_summary(self):
        self.doc.add_heading("1. 요약 (Executive Summary)", level=1)
        s = self.stats

        p = self.doc.add_paragraph()
        p.add_run("RFP 대상 법령 (6개 유형): ").bold = True
        p.add_run("UU, PP, PERPRES, PERPPU, PERMEN, PERBAN")

        self.doc.add_paragraph()
        self.doc.add_heading("핵심 수치", level=2)

        summary = [
            ["단계", "건수", "비율", "비고"],
            ["① 총 법령", fmt(s["total"]), "100%", f"{s['year_min']}~{s['year_max']}년"],
            ["② PDF URL 있음", fmt(s["pdf_url_exists"]), pct(s["pdf_url_exists"], s["total"]), ""],
            ["   └ PDF URL 없음", fmt(s["pdf_url_none"]), pct(s["pdf_url_none"], s["total"]), "원본에 PDF 없음"],
            ["③ PDF 다운로드", fmt(s["pdf_downloaded"]), pct(s["pdf_downloaded"], s["pdf_url_exists"]), ""],
            ["   └ 다운로드 실패", fmt(s["pdf_download_failed"]), "", ""],
            ["④ 텍스트 추출", fmt(s["extract_success"]), pct(s["extract_success"], s["pdf_downloaded"]), ""],
            ["   └ OCR 필요", fmt(s["extract_needs_ocr"]), "", "스캔 이미지"],
            ["⑤ 구조 파싱", fmt(s["parse_success"]), pct(s["parse_success"], s["extract_success"]), ""],
        ]
        create_table(self.doc, summary)

        self.doc.add_paragraph()
        self.doc.add_heading("추출 데이터 규모", level=2)

        ref = s.get("references", {})
        data_scale = [
            ["항목", "수치"],
            ["총 텍스트 용량", f"{s['total_chars'] / 1024 / 1024:.0f} MB"],
            ["총 조항 (Pasal)", fmt(s["total_pasal"])],
            ["총 항 (Ayat)", fmt(s["total_ayat"])],
            ["외부 참조 관계", fmt(ref.get("total_external", 0))],
            ["내부 참조 관계", fmt(ref.get("total_internal", 0))],
        ]
        create_table(self.doc, data_scale)

        self.doc.add_page_break()

    def add_collection_section(self):
        self.doc.add_heading("2. 데이터 수집 현황", level=1)
        s = self.stats

        self.doc.add_heading("유형별 현황", level=2)
        type_data = [["약어", "법령 유형", "한국어", "건수", "비율"]]
        for jenis, count in s["by_type"]:
            abbrev = RFP_TYPE_ABBREV.get(jenis, "")
            korean = RFP_TYPE_KOREAN.get(jenis, "")
            type_data.append([abbrev, jenis, korean, fmt(count), pct(count, s["total"])])
        create_table(self.doc, type_data)

        self.doc.add_page_break()

    def add_pdf_section(self):
        self.doc.add_heading("3. PDF 파일 현황", level=1)
        s = self.stats

        self.doc.add_heading("3.1 수집 흐름", level=2)
        flow = [
            ["단계", "건수", "비율"],
            ["총 법령", fmt(s["total"]), "100%"],
            ["├ PDF URL 있음", fmt(s["pdf_url_exists"]), pct(s["pdf_url_exists"], s["total"])],
            ["│  ├ 다운로드 성공", fmt(s["pdf_downloaded"]), pct(s["pdf_downloaded"], s["pdf_url_exists"])],
            ["│  └ 다운로드 실패", fmt(s["pdf_download_failed"]), ""],
            ["└ PDF URL 없음", fmt(s["pdf_url_none"]), pct(s["pdf_url_none"], s["total"])],
        ]
        create_table(self.doc, flow)

        self.doc.add_paragraph()
        self.doc.add_heading("3.2 유형별 PDF 현황", level=2)
        type_pdf = [["약어", "총", "다운로드", "확보율"]]
        for jenis, total, url, downloaded, extracted in s["type_pdf"]:
            abbrev = RFP_TYPE_ABBREV.get(jenis, "")
            type_pdf.append([abbrev, fmt(total), fmt(downloaded), pct(downloaded, total)])
        create_table(self.doc, type_pdf)

        self.doc.add_page_break()

    def add_extraction_section(self):
        self.doc.add_heading("4. 텍스트 추출 현황", level=1)
        s = self.stats

        self.doc.add_heading("4.1 추출 결과", level=2)
        extract = [
            ["단계", "건수", "비율"],
            ["PDF 보유", fmt(s["pdf_downloaded"]), "100%"],
            ["├ 추출 성공", fmt(s["extract_success"]), pct(s["extract_success"], s["pdf_downloaded"])],
            ["├ OCR 필요", fmt(s["extract_needs_ocr"]), pct(s["extract_needs_ocr"], s["pdf_downloaded"])],
            ["└ 실패/미처리", fmt(s["extract_failed"]), ""],
        ]
        create_table(self.doc, extract)

        self.doc.add_paragraph()
        self.doc.add_heading("4.2 문서 길이 분포", level=2)
        length_data = [["문서 길이", "건수"]]
        for length_group, cnt in s["length_dist"]:
            length_data.append([length_group, fmt(cnt)])
        create_table(self.doc, length_data)

        self.doc.add_page_break()

    def add_parsing_section(self):
        self.doc.add_heading("5. 문서 구조 파싱", level=1)
        s = self.stats

        self.doc.add_heading("5.1 파싱 결과", level=2)
        parsing = [
            ["구조", "인도네시아어", "추출 건수"],
            ["장(章)", "BAB", fmt(s["total_bab"])],
            ["조(條)", "Pasal", fmt(s["total_pasal"])],
            ["항(項)", "Ayat", fmt(s["total_ayat"])],
            ["호(號)", "Huruf", fmt(s["total_huruf"])],
        ]
        create_table(self.doc, parsing)

        self.doc.add_paragraph()
        self.doc.add_heading("5.2 유형별 평균 조항 수", level=2)
        avg_data = [["약어", "법령수", "평균 조", "평균 항", "최대 조"]]
        for jenis, cnt, avg_pasal, avg_ayat, max_pasal in s["avg_pasal_by_type"]:
            abbrev = RFP_TYPE_ABBREV.get(jenis, "")
            avg_data.append([abbrev, fmt(cnt), str(avg_pasal), str(avg_ayat), fmt(max_pasal)])
        create_table(self.doc, avg_data)

        self.doc.add_paragraph()
        self.doc.add_heading("5.3 가장 긴 법령 TOP 5", level=2)
        longest = [["유형", "번호", "연도", "제목", "조 수"]]
        for jenis, nomor, tahun, tentang, pasal_cnt in s["longest_docs"]:
            abbrev = RFP_TYPE_ABBREV.get(jenis, "")
            longest.append([abbrev, str(nomor), str(tahun), tentang + "...", fmt(pasal_cnt)])
        create_table(self.doc, longest)

        self.doc.add_page_break()

    def add_reference_section(self):
        self.doc.add_heading("6. 참조/인용 분석", level=1)
        ref = self.stats.get("references", {})

        if not ref:
            self.doc.add_paragraph("참조 데이터 없음")
            self.doc.add_page_break()
            return

        self.doc.add_heading("6.1 참조 현황", level=2)
        ref_summary = [
            ["항목", "건수"],
            ["분석 대상 법령", fmt(ref.get("total_docs", 0))],
            ["외부 참조 (타 법령 인용)", fmt(ref.get("total_external", 0))],
            ["내부 참조 (자체 조항 참조)", fmt(ref.get("total_internal", 0))],
            ["조건부 조항", fmt(ref.get("total_conditional", 0))],
        ]
        create_table(self.doc, ref_summary)

        self.doc.add_paragraph()
        self.doc.add_heading("6.2 가장 많이 인용된 법령 TOP 10", level=2)
        most_cited = [["인용 횟수", "법령"]]
        for cited, cnt in ref.get("most_cited", []):
            most_cited.append([fmt(cnt), cited])
        create_table(self.doc, most_cited)

        self.doc.add_page_break()

    def add_agency_year_section(self):
        self.doc.add_heading("7. 발급기관 및 연도별 분석", level=1)
        s = self.stats

        self.doc.add_heading("7.1 발급기관별 법령 수 TOP 10", level=2)
        agency_data = [["발급기관", "건수"]]
        for agency, cnt in s["by_agency"]:
            agency_data.append([agency, fmt(cnt)])
        create_table(self.doc, agency_data)

        self.doc.add_paragraph()
        self.doc.add_heading("7.2 연도별 법령 생산량 TOP 10", level=2)
        year_data = [["연도", "건수"]]
        for year, cnt in s["by_year"]:
            year_data.append([str(year), fmt(cnt)])
        create_table(self.doc, year_data)

        self.doc.add_page_break()

    def add_validity_section(self):
        self.doc.add_heading("8. 현행성 판단", level=1)
        s = self.stats

        validity = [
            ["효력 상태", "건수", "비율"],
            ["현행 (Berlaku)", fmt(s["status_berlaku"]), pct(s["status_berlaku"], s["total"])],
            ["폐지 (Tidak Berlaku)", fmt(s["status_tidak"]), pct(s["status_tidak"], s["total"])],
            ["기타", fmt(s["total"] - s["status_berlaku"] - s["status_tidak"]), ""],
        ]
        create_table(self.doc, validity)

        self.doc.add_page_break()

    def add_issues_section(self):
        self.doc.add_heading("9. 문제점 및 권고사항", level=1)
        s = self.stats

        self.doc.add_heading("9.1 문제점", level=2)
        issues = [
            ["문제", "건수", "설명"],
            ["PDF 미확보", fmt(s["pdf_url_none"]), "원본 사이트에 PDF 없음"],
            ["다운로드 실패", fmt(s["pdf_download_failed"]), "서버 오류 등"],
            ["OCR 필요", fmt(s["extract_needs_ocr"]), "스캔 이미지 PDF"],
            ["추출 실패", fmt(s["extract_failed"]), "파일 손상 등"],
        ]
        create_table(self.doc, issues)

        self.doc.add_paragraph()
        self.doc.add_heading("9.2 권고사항", level=2)
        recommendations = [
            f"1. OCR 처리: 스캔 이미지 PDF {fmt(s['extract_needs_ocr'])}건 처리 필요",
            f"2. 다운로드 재시도: 실패 {fmt(s['pdf_download_failed'])}건 재시도",
            "3. 구법령 보완: 1945-1980년대 PDF 미확보 건 타 기관 협조",
        ]
        for rec in recommendations:
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run(rec)

        self.doc.add_page_break()

    def add_appendix(self):
        self.doc.add_heading("부록: 용어 대조표", level=1)

        self.doc.add_heading("RFP 대상 법령 유형", level=2)
        terms = [
            ["약어", "인도네시아어", "한국어"],
            ["UU", "Undang-Undang", "법률"],
            ["PP", "Peraturan Pemerintah", "정부령"],
            ["PERPRES", "Peraturan Presiden", "대통령령"],
            ["PERPPU", "Peraturan Pemerintah Pengganti UU", "긴급법률"],
            ["PERMEN", "Peraturan Menteri", "장관령"],
            ["PERBAN", "Peraturan Badan/Lembaga", "기관규정"],
        ]
        create_table(self.doc, terms)

        self.doc.add_paragraph()
        self.doc.add_heading("문서 구조", level=2)
        structure = [
            ["인도네시아어", "한국어"],
            ["BAB", "장(章)"],
            ["Pasal", "조(條)"],
            ["Ayat", "항(項)"],
            ["Huruf", "호(號)"],
        ]
        create_table(self.doc, structure)

    def generate(self):
        print("보고서 생성 중...")

        self.add_cover()
        self.add_toc()
        self.add_summary()
        self.add_collection_section()
        self.add_pdf_section()
        self.add_extraction_section()
        self.add_parsing_section()
        self.add_reference_section()
        self.add_agency_year_section()
        self.add_validity_section()
        self.add_issues_section()
        self.add_appendix()

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        self.doc.save(OUTPUT_FILE)
        print(f"저장 완료: {OUTPUT_FILE}")
        return OUTPUT_FILE


def main():
    print("=" * 50)
    print("인도네시아 법령정보시스템 - 종합분석 보고서")
    print("=" * 50)
    ReportGenerator().generate()


if __name__ == "__main__":
    main()
